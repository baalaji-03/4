from io import BytesIO
from datetime import date
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

HEADER_FILL = PatternFill(start_color="0F766E", end_color="0F766E", fill_type="solid")
HEADER_FONT = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
ALERT_RED_FILL = PatternFill(start_color="FEE2E2", end_color="FEE2E2", fill_type="solid")
ALERT_ORANGE_FILL = PatternFill(start_color="FFEDD5", end_color="FFEDD5", fill_type="solid")
THIN_BORDER = Border(
    left=Side(style='thin', color='CBD5E1'),
    right=Side(style='thin', color='CBD5E1'),
    top=Side(style='thin', color='CBD5E1'),
    bottom=Side(style='thin', color='CBD5E1')
)

def _style_header(ws, col_count):
    ws.row_dimensions[1].height = 24
    for col in range(1, col_count + 1):
        cell = ws.cell(row=1, column=col)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

def _autofit_columns(ws):
    for col in ws.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

# ----------------- Export Functions -----------------
def export_inventory_excel(medicines: list) -> BytesIO:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Current Inventory"

    headers = [
        "Item Code", "Medicine Name", "Generic Name", "Manufacturer",
        "Category", "HSN Code", "GST %", "Rack Location", "Unit",
        "Min Stock", "Total Available Stock", "Default Pur. Price", "Default Sell Price", "Status"
    ]
    ws.append(headers)

    for med in medicines:
        total_stock = sum(b.quantity for b in med.batches if b.status != "EXPIRED")
        status_text = "LOW STOCK" if total_stock <= med.min_stock_level else "IN STOCK"
        ws.append([
            med.item_code or "",
            med.name,
            med.generic_name or "",
            med.manufacturer or "",
            med.category or "",
            med.hsn_code or "3004",
            float(med.gst_rate or 12.0),
            med.rack_location or "",
            med.unit or "Strip",
            med.min_stock_level,
            total_stock,
            float(med.default_purchase_price or 0.0),
            float(med.default_selling_price or 0.0),
            status_text
        ])

    _style_header(ws, len(headers))
    _autofit_columns(ws)
    
    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf

def export_batches_and_expiry_excel(batches: list) -> BytesIO:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Batches & Expiry"

    headers = [
        "Medicine Name", "Batch Number", "Mfg Date", "Expiry Date",
        "Days Remaining", "Quantity", "Purchase Price", "MRP", "Selling Price",
        "Total Stock Value", "Status"
    ]
    ws.append(headers)

    today = date.today()
    for row_idx, b in enumerate(batches, start=2):
        days_rem = (b.expiry_date - today).days if b.expiry_date else 0
        stock_val = float(b.purchase_price or 0) * b.quantity
        
        status = "EXPIRED" if days_rem <= 0 else ("NEAR EXPIRY" if days_rem <= 180 else "AVAILABLE")
        ws.append([
            b.medicine.name if b.medicine else "",
            b.batch_number,
            b.manufacturing_date.strftime("%d-%m-%Y") if b.manufacturing_date else "",
            b.expiry_date.strftime("%d-%m-%Y") if b.expiry_date else "",
            days_rem,
            b.quantity,
            float(b.purchase_price or 0),
            float(b.mrp or 0),
            float(b.selling_price or 0),
            stock_val,
            status
        ])

        # Color row if expired or near expiry
        if status == "EXPIRED":
            for col in range(1, len(headers) + 1):
                ws.cell(row=row_idx, column=col).fill = ALERT_RED_FILL
        elif status == "NEAR EXPIRY":
            for col in range(1, len(headers) + 1):
                ws.cell(row=row_idx, column=col).fill = ALERT_ORANGE_FILL

    _style_header(ws, len(headers))
    _autofit_columns(ws)

    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf

def export_sales_excel(sales: list) -> BytesIO:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sales Report"

    headers = [
        "Invoice No", "Date", "Customer Name", "Customer Mobile",
        "Payment Method", "Subtotal", "Discount", "Taxable Amount",
        "CGST", "SGST", "IGST", "Total GST", "Round Off", "Net Amount", "Gross Profit"
    ]
    ws.append(headers)

    for s in sales:
        ws.append([
            s.invoice_number,
            s.invoice_date.strftime("%d-%m-%Y") if s.invoice_date else "",
            s.customer_name or "Walk-in",
            s.customer_mobile or "",
            s.payment_method,
            float(s.subtotal or 0),
            float(s.discount_amount or 0),
            float(s.taxable_amount or 0),
            float(s.cgst_amount or 0),
            float(s.sgst_amount or 0),
            float(s.igst_amount or 0),
            float(s.tax_amount or 0),
            float(s.round_off or 0),
            float(s.net_amount or 0),
            float(s.gross_profit or 0)
        ])

    _style_header(ws, len(headers))
    _autofit_columns(ws)

    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf

def export_gst_summary_excel(sales: list) -> BytesIO:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "GSTR-1 Summary"

    headers = [
        "Invoice No", "Invoice Date", "Customer / Party", "GSTIN",
        "Place of Supply", "Invoice Value", "Taxable Value", "GST Rate %",
        "CGST Amount", "SGST Amount", "IGST Amount", "Total Tax"
    ]
    ws.append(headers)

    for s in sales:
        ws.append([
            s.invoice_number,
            s.invoice_date.strftime("%d-%m-%Y") if s.invoice_date else "",
            s.customer_name or "B2C Consumer",
            "URP", # Unregistered person
            "33 - Tamil Nadu",
            float(s.net_amount or 0),
            float(s.taxable_amount or 0),
            "12%",
            float(s.cgst_amount or 0),
            float(s.sgst_amount or 0),
            float(s.igst_amount or 0),
            float(s.tax_amount or 0)
        ])

    _style_header(ws, len(headers))
    _autofit_columns(ws)

    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf

# ----------------- Excel Templates -----------------
def generate_medicine_template() -> BytesIO:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Medicines Import"

    headers = [
        "Item Code", "Medicine Name *", "Generic Name", "Manufacturer",
        "Category", "HSN Code", "GST Rate %", "Rack Location", "Unit",
        "Min Stock", "Default Pur Price", "Default Sell Price", "Barcode"
    ]
    ws.append(headers)
    ws.append([
        "MED-1001", "Paracetamol 500mg", "Acetaminophen", "Apex Pharma",
        "Analgesic", "3004", 12.0, "Rack-A2", "Strip", 20, 15.00, 25.00, "8901234567890"
    ])

    _style_header(ws, len(headers))
    _autofit_columns(ws)

    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf

def generate_purchase_template() -> BytesIO:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Purchase Import"

    headers = [
        "Supplier Invoice No *", "Invoice Date (YYYY-MM-DD) *", "Supplier Code or Name *",
        "Medicine Name *", "Batch Number *", "Expiry Date (YYYY-MM-DD) *",
        "Quantity *", "Free Quantity", "Purchase Price *", "MRP *", "Selling Price *",
        "Discount %", "GST Rate %"
    ]
    ws.append(headers)
    ws.append([
        "INV-2026-901", "2026-10-15", "Apollo Med Wholesale Ltd",
        "Dolo 650mg Tablet", "DL8890", "2028-05-31",
        100, 10, 24.50, 33.60, 33.60, 2.0, 12.0
    ])

    _style_header(ws, len(headers))
    _autofit_columns(ws)

    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf

# ----------------- Excel Upload Parsing & Validation -----------------
def parse_and_validate_medicine_excel(file_bytes: bytes) -> dict:
    wb = openpyxl.load_workbook(BytesIO(file_bytes), data_only=True)
    ws = wb.active

    rows = list(ws.iter_rows(values_only=True))
    if not rows or len(rows) < 2:
        return {"success": False, "error": "Excel file is empty or missing data rows", "rows": []}

    parsed = []
    errors = []

    for idx, r in enumerate(rows[1:], start=2):
        if not r or not any(r):
            continue
        med_name = str(r[1]).strip() if len(r) > 1 and r[1] is not None else ""
        if not med_name:
            errors.append(f"Row {idx}: Medicine Name is required.")
            continue

        try:
            gst = float(r[6]) if len(r) > 6 and r[6] is not None else 12.0
            min_stk = int(r[9]) if len(r) > 9 and r[9] is not None else 10
            p_price = float(r[10]) if len(r) > 10 and r[10] is not None else 0.0
            s_price = float(r[11]) if len(r) > 11 and r[11] is not None else 0.0
        except ValueError as e:
            errors.append(f"Row {idx}: Numeric format error ({e}).")
            continue

        parsed.append({
            "row_index": idx,
            "item_code": str(r[0]).strip() if len(r) > 0 and r[0] else "",
            "name": med_name,
            "generic_name": str(r[2]).strip() if len(r) > 2 and r[2] else "",
            "manufacturer": str(r[3]).strip() if len(r) > 3 and r[3] else "",
            "category": str(r[4]).strip() if len(r) > 4 and r[4] else "Allopathy",
            "hsn_code": str(r[5]).strip() if len(r) > 5 and r[5] else "3004",
            "gst_rate": gst,
            "rack_location": str(r[7]).strip() if len(r) > 7 and r[7] else "",
            "unit": str(r[8]).strip() if len(r) > 8 and r[8] else "Strip",
            "min_stock_level": min_stk,
            "default_purchase_price": p_price,
            "default_selling_price": s_price,
            "barcode": str(r[12]).strip() if len(r) > 12 and r[12] else None
        })

    return {
        "success": len(errors) == 0 or len(parsed) > 0,
        "total_rows": len(parsed),
        "errors": errors,
        "preview": parsed[:10],
        "all_data": parsed
    }
